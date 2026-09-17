# -*- coding: utf-8 -*-
"""
RELATÓRIO SEMANAL — gerador de referência (RV Engenharia)

Uso:
    python relatorio_semanal.py --dados dados/obra_selecta.json --semana 8 --saida ./saida

Produz, para a semana pedida:
    Resumo_Semanal_<obra>_<inicio>-a-<fim>.pdf         (1 página, A4 paisagem)
    Controle_Diarias_<obra>_Semana<N>_<inicio>-a-<fim>.xlsx  (Parâmetros + 1 aba por dia + Resumo Semanal)

Este script É o padrão visual. Ao portar para outra linguagem/stack, replique
as medidas, cores, ordem de blocos e textos daqui. Não invente layout.
"""
import argparse
import datetime as dt
import json
import os
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.colors import HexColor, white

# ============================================================ IDENTIDADE VISUAL
COR = dict(
    azul="#1F3864", azul_claro="#D9E1F2", cinza="#F2F2F2", linha="#BFBFBF",
    verde="#C6E0B4", texto="#262626", vermelho="#C00000", cinza_txt="#595959",
    rodape="#808080", input_xlsx="FFF2CC",
)
FONTE = "Helvetica"
FONTE_XLSX = "Arial"
DIAS_SEMANA = ["SEGUNDA-FEIRA", "TERÇA-FEIRA", "QUARTA-FEIRA", "QUINTA-FEIRA",
               "SEXTA-FEIRA", "SÁBADO", "DOMINGO"]
DIAS_CURTO = ["Segunda", "Terça", "Quarta", "Quinta", "Sexta", "Sábado", "Domingo"]


# ================================================================ UTILIDADES
def brl(v):
    s = f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"R$ {s}"


def num(v):
    return f"{v:g}".replace(".", ",")


def d(s):
    return dt.date.fromisoformat(s)


def carregar(caminho):
    with open(caminho, encoding="utf-8") as f:
        j = json.load(f)
    j["_equipe"] = OrderedDict((e["nome"], e) for e in j["equipe"])
    for l in j["lancamentos"]:
        l["_data"] = d(l["data"])
    j["lancamentos"].sort(key=lambda x: x["_data"])
    for s in j["semanas"]:
        s["_inicio"], s["_fim"] = d(s["inicio"]), d(s["fim"])
    return j


def extras_de(x):
    return [(e["descricao"], e.get("qtd", 1), float(e["valor"])) for e in x.get("extras", [])]


def bonus_de(x, equipe):
    """Bônus de produção do dia: lista (nome, diárias_bonus, valor)."""
    b = x.get("bonus")
    if not b:
        return []
    return [(n, q, equipe[n]["diaria"] * q) for n, q in b["diarias"].items()]


def stats(dias, equipe):
    """alim = quentinhas + extras (entregas, gelo etc.). mo inclui bônus de produção."""
    diarias = sum(sum(x["presencas"].values()) for x in dias)
    bonus = sum(v for x in dias for _, _, v in bonus_de(x, equipe))
    bonus_d = sum(q for x in dias for _, q, _ in bonus_de(x, equipe))
    mo = sum(sum(equipe[n]["diaria"] * f for n, f in x["presencas"].items()) for x in dias) + bonus
    quent = sum(x["quentinhas"] for x in dias)
    quent_custo = sum(x["quentinhas"] * x["valor_quentinha"] for x in dias)
    extras = sum(q * v for x in dias for _, q, v in extras_de(x))
    alim = quent_custo + extras
    return dict(dias=len(dias), diarias=diarias, mo=mo, bonus=bonus, bonus_d=bonus_d, quent=quent,
                quent_custo=quent_custo, extras=extras, alim=alim, total=mo + alim)


def extras_agrupados(dias):
    g = OrderedDict()
    for x in dias:
        for desc, q, v in extras_de(x):
            g.setdefault(desc, {"qtd": 0, "custo": 0.0, "dias": [], "valor": v, "misto": False})
            if g[desc]["valor"] != v:
                g[desc]["misto"] = True
            g[desc]["qtd"] += q
            g[desc]["custo"] += q * v
            g[desc]["dias"].append(x["_data"].strftime("%d/%m"))
    return g


def qtd_txt(g):
    """'8 × R$ 6,00' quando o unitário é fixo; '3 un.' quando variou no período."""
    return f"{g['qtd']} un." if g["misto"] else f"{g['qtd']} × {brl(g['valor'])}"


def semana_info(j, numero):
    sem = next(s for s in j["semanas"] if s["numero"] == numero)
    dias = [x for x in j["lancamentos"] if sem["_inicio"] <= x["_data"] <= sem["_fim"]]
    anteriores = []
    for s in j["semanas"]:
        if s["numero"] < numero:
            ds = [x for x in j["lancamentos"] if s["_inicio"] <= x["_data"] <= s["_fim"]]
            if not ds:
                continue
            s = dict(s, _inicio=ds[0]["_data"], _fim=ds[-1]["_data"])   # período efetivamente lançado
            anteriores.append((s, stats(ds, j["_equipe"])))
    return sem, dias, anteriores


def equipe_da_semana(j, dias):
    """Só quem teve pelo menos uma presença na semana, na ordem do cadastro."""
    presentes = set()
    for x in dias:
        presentes.update(x["presencas"].keys())
        if x.get("bonus"):
            presentes.update(x["bonus"]["diarias"].keys())
    return [e for n, e in j["_equipe"].items() if n in presentes]


def faixas_quentinha(dias):
    f = OrderedDict()
    for x in dias:
        if x["quentinhas"] == 0:
            continue
        f.setdefault(x["valor_quentinha"], {"qtd": 0, "dias": []})
        f[x["valor_quentinha"]]["qtd"] += x["quentinhas"]
        f[x["valor_quentinha"]]["dias"].append(x["_data"].strftime("%d/%m"))
    return OrderedDict(sorted(f.items()))


def historico_precos(j, ate):
    hist = []
    for x in j["lancamentos"]:
        if x["_data"] > ate or x["quentinhas"] == 0:
            continue
        if hist and hist[-1]["valor"] == x["valor_quentinha"]:
            hist[-1]["fim"] = x["_data"]
        else:
            hist.append(dict(valor=x["valor_quentinha"], inicio=x["_data"], fim=x["_data"]))
    return hist


# ====================================================================== PDF
def gerar_pdf(j, numero, saida, parcial=False):
    sem, dias, anteriores = semana_info(j, numero)
    equipe = j["_equipe"]
    eq_sem = equipe_da_semana(j, dias)
    st = stats(dias, equipe)
    faixas = faixas_quentinha(dias)
    ext_sem = extras_agrupados(dias)
    obra = j["obra"]
    ini, fim = dias[0]["_data"], dias[-1]["_data"]
    W, H = landscape(A4)
    M = 26
    C = {k: HexColor(v) for k, v in COR.items() if v.startswith("#")}
    B, R, O = FONTE + "-Bold", FONTE, FONTE + "-Oblique"

    # ---- escala vertical: tudo tem de caber em UMA página
    n_func, n_sem = len(eq_sem), len(anteriores) + 1
    max_pres = max(len(x["presencas"]) for x in dias)
    max_extras = max(len(extras_de(x)) + (1 if x.get("bonus") else 0) for x in dias)
    linhas_grade = 2 if len(dias) > 3 else 1
    GAP = 12
    fixo = 34 + 10 + 14 + 15 + 8 + 26 + (linhas_grade - 1) * GAP
    variavel = (linhas_grade * (16 + 12 + 12 * max_pres + 13 + 13 + 13 * max_extras + 10 + 17)
                + (13 + 11 + 11 * n_func + 14) + (13 + 10 + 10 * n_sem + 13))
    k = min(1.0, (H - 2 * M - fixo) / variavel)
    fk = max(0.78, k)

    def F(v):
        return round(v * fk, 1)

    RD, RC, RO = 12 * k, 13 * k, 10 * k
    RF, RA, RCARD = 11 * k, 10 * k, 22 * k
    CW = (W - 2 * M - 2 * GAP) / 3
    BLH = 13 + 11 + RF * n_func + 14
    ACH = 13 + 10 + RA * n_sem + 13
    necessario = 16 + 12 + RD * max_pres + RC * 2 + RC * max_extras + RO + 17
    disponivel = H - 2 * M - fixo - BLH - ACH
    BH = max(necessario, min(disponivel / linhas_grade, 150)) if linhas_grade == 2 else max(necessario, 118)

    nome_arq = (f"Resumo_Semanal_{obra['nome'].split('–')[0].strip().replace(' ', '_')}_"
                f"{ini.strftime('%d-%m')}-a-{fim.strftime('%d-%m')}.pdf")
    c = canvas.Canvas(os.path.join(saida, nome_arq), pagesize=landscape(A4))
    c.setTitle(f"Resumo Semanal de Diárias — {obra['nome']} — {ini.strftime('%d/%m')} a {fim.strftime('%d/%m/%Y')}")

    def linha(x1, y, x2):
        c.setStrokeColor(C["linha"])
        c.setLineWidth(0.3)
        c.line(x1, y, x2, y)

    def caixa(x, y, w, h, lw=0.8):
        c.setStrokeColor(C["azul"])
        c.setLineWidth(lw)
        c.rect(x, y, w, h, stroke=1, fill=0)

    def faixa(x, y, w, h, cor, txt, tam, cor_txt=white, direita=None, tam_d=None):
        c.setFillColor(cor)
        c.rect(x, y - h, w, h, stroke=0, fill=1)
        c.setFillColor(cor_txt)
        c.setFont(B, tam)
        c.drawString(x + 6, y - h + (h - tam) / 2 + 1, txt)
        if direita:
            c.setFont(R, tam_d or tam)
            c.drawRightString(x + w - 6, y - h + (h - tam) / 2 + 1, direita)

    # ---- cabeçalho
    y = H - M
    c.setFillColor(C["azul"])
    c.rect(M, y - 34, W - 2 * M, 34, stroke=0, fill=1)
    c.setFillColor(white)
    c.setFont(B, 14)
    c.drawString(M + 10, y - 15, f"CONTROLE DE DIÁRIAS E CUSTOS DE MÃO DE OBRA — SEMANA "
                                 f"{ini.strftime('%d/%m')} A {fim.strftime('%d/%m/%Y')}"
                                 + ("  ·  PARCIAL" if parcial else ""))
    c.setFont(R, 8.5)
    c.drawString(M + 10, y - 27, f"Obra: {obra['nome']}  ·  {obra['endereco']}")
    c.setFont(B, 10)
    c.drawRightString(W - M - 10, y - 17, obra["empresa"].upper())
    top = y - 34 - 10

    # ---- blocos de dia
    def bloco_dia(x, ytop, dia):
        data = dia["_data"]
        pres = [(n, dia["presencas"][n]) for n in equipe if n in dia["presencas"]]
        mo = sum(equipe[n]["diaria"] * f for n, f in pres)
        q_custo = dia["quentinhas"] * dia["valor_quentinha"]
        ext = extras_de(dia)
        bon = bonus_de(dia, equipe)
        bon_v = sum(v for _, _, v in bon)
        tot = mo + bon_v + q_custo + sum(q * v for _, q, v in ext)
        diarias = sum(f for _, f in pres)
        caixa(x, ytop - BH, CW, BH)
        faixa(x, ytop, CW, 16, C["azul"], DIAS_SEMANA[data.weekday()], 9,
              direita=data.strftime("%d/%m/%Y"), tam_d=8)
        yy = ytop - 16
        c.setFillColor(C["azul_claro"])
        c.rect(x, yy - 12, CW, 12, stroke=0, fill=1)
        c.setFillColor(C["azul"])
        c.setFont(B, 7)
        c.drawString(x + 5, yy - 8.5, "FUNCIONÁRIO")
        c.drawString(x + 0.42 * CW, yy - 8.5, "FUNÇÃO")
        c.drawCentredString(x + 0.74 * CW, yy - 8.5, "DIÁRIAS")
        c.drawRightString(x + CW - 5, yy - 8.5, "VALOR")
        yy -= 12
        for n, f in pres:
            c.setFillColor(C["texto"])
            c.setFont(R, F(7.5))
            c.drawString(x + 5, yy - RD + 3.5 * k, n)
            c.setFont(R, F(7))
            c.drawString(x + 0.42 * CW, yy - RD + 3.5 * k, equipe[n]["funcao"])
            c.setFont(R, F(7.5))
            c.drawCentredString(x + 0.74 * CW, yy - RD + 3.5 * k, num(f))
            c.drawRightString(x + CW - 5, yy - RD + 3.5 * k, brl(equipe[n]["diaria"] * f))
            linha(x + 4, yy - RD + 0.5, x + CW - 4)
            yy -= RD
        itens = [(f"Subtotal mão de obra ({num(diarias)} diária{'s' if diarias != 1 else ''})", brl(mo))]
        if bon:
            itens.append((f"Bônus de produção: {len(bon)} pessoas × {num(bon[0][1])} diárias", brl(bon_v)))
        itens.append((f"Quentinhas: {dia['quentinhas']} un. × {brl(dia['valor_quentinha'])}", brl(q_custo)))
        for desc, q, v in ext:
            itens.append((f"{desc}: {q} × {brl(v)}" if q != 1 else desc, brl(q * v)))
        for lab, val in itens:
            c.setFillColor(C["cinza"])
            c.rect(x, yy - RC, CW, RC, stroke=0, fill=1)
            c.setFillColor(C["texto"])
            c.setFont(B, F(7.5))
            c.drawString(x + 5, yy - RC + 4 * k, lab)
            c.drawRightString(x + CW - 5, yy - RC + 4 * k, val)
            yy -= RC
        if dia.get("observacao"):
            c.setFillColor(C["vermelho"])
            c.setFont(O, F(6.2))
            txt = dia["observacao"]
            while c.stringWidth(txt, O, F(6.2)) > CW - 12 and " " in txt:
                txt = txt.rsplit(" ", 1)[0]
            c.drawString(x + 5, yy - 8 * k, txt + ("…" if txt != dia["observacao"] else ""))
        c.setFillColor(C["verde"])
        c.rect(x, ytop - BH, CW, 17, stroke=0, fill=1)
        caixa(x, ytop - BH, CW, 17)
        c.setFillColor(C["azul"])
        c.setFont(B, 9)
        c.drawString(x + 6, ytop - BH + 5.5, "TOTAL DO DIA")
        c.drawRightString(x + CW - 6, ytop - BH + 5.5, brl(tot))

    for i, dia in enumerate(dias):
        col, row = i % 3, i // 3
        bloco_dia(M + col * (CW + GAP), top - row * (BH + GAP), dia)

    # ---- card de observações nos espaços vagos da grade
    vagos = 3 * linhas_grade - len(dias)
    if vagos > 0:
        i = len(dias)
        col, row = i % 3, i // 3
        x6, y6 = M + col * (CW + GAP), top - row * (BH + GAP)
        CW6 = vagos * CW + (vagos - 1) * GAP
        caixa(x6, y6 - BH, CW6, BH)
        faixa(x6, y6, CW6, 16, C["azul"], "OBSERVAÇÕES DA SEMANA", 9)
        yy = y6 - 16
        notas = [(x["_data"].strftime("%d/%m"), x["observacao"]) for x in dias if x.get("observacao")]
        for dd, t in notas[: int((BH - 16 - 17) // RCARD)]:
            c.setFillColor(C["azul"])
            c.setFont(B, F(7))
            c.drawString(x6 + 6, yy - 10 * k, dd)
            c.setFillColor(C["texto"])
            c.setFont(R, F(6.6))
            txt = t
            while c.stringWidth(txt, R, F(6.6)) > CW6 - 44 and " " in txt:
                txt = txt.rsplit(" ", 1)[0]
            c.drawString(x6 + 34, yy - 10 * k, txt + ("…" if txt != t else ""))
            linha(x6 + 5, yy - RCARD + 6 * k, x6 + CW6 - 5)
            yy -= RCARD

    # ---- faixa do resumo
    ry = top - linhas_grade * BH - (linhas_grade - 1) * GAP - 14
    faixa(M, ry, W - 2 * M, 15, C["azul"], "RESUMO GERAL DA SEMANA", 10)
    ry -= 15
    LW = (W - 2 * M) * 0.60
    RW = (W - 2 * M) - LW - GAP
    RX = M + LW + GAP

    # ---- esquerda: gastos por funcionário
    diarias_f = {e["nome"]: 0 for e in eq_sem}
    bonus_f = {e["nome"]: 0 for e in eq_sem}
    for x in dias:
        for n, f in x["presencas"].items():
            diarias_f[n] += f
        for n, q, _ in bonus_de(x, equipe):
            bonus_f[n] += q
    caixa(M, ry - BLH, LW, BLH)
    faixa(M, ry, LW, 13, C["azul_claro"], "GASTOS POR FUNCIONÁRIO", 8, cor_txt=C["azul"])
    yy = ry - 13
    c.setFillColor(C["cinza"])
    c.rect(M, yy - 11, LW, 11, stroke=0, fill=1)
    c.setFillColor(C["azul"])
    c.setFont(B, 7)
    c.drawString(M + 6, yy - 7.8, "FUNCIONÁRIO")
    c.drawString(M + 0.30 * LW, yy - 7.8, "FUNÇÃO")
    c.drawCentredString(M + 0.55 * LW, yy - 7.8, "VALOR DA DIÁRIA")
    c.drawCentredString(M + 0.74 * LW, yy - 7.8, "DIÁRIAS NA SEMANA")
    c.drawRightString(M + LW - 6, yy - 7.8, "TOTAL A PAGAR")
    yy -= 11
    for e in eq_sem:
        n = e["nome"]
        c.setFillColor(C["texto"])
        c.setFont(R, F(7.5))
        c.drawString(M + 6, yy - RF + 3 * k, n)
        c.setFont(R, F(7))
        c.drawString(M + 0.30 * LW, yy - RF + 3 * k, e["funcao"])
        c.setFont(R, F(7.5))
        c.drawCentredString(M + 0.55 * LW, yy - RF + 3 * k, brl(e["diaria"]))
        c.drawCentredString(M + 0.74 * LW, yy - RF + 3 * k,
                            num(diarias_f[n]) + (f" + {num(bonus_f[n])} bônus" if bonus_f[n] else ""))
        c.setFont(B, F(7.5))
        c.drawRightString(M + LW - 6, yy - RF + 3 * k, brl(e["diaria"] * (diarias_f[n] + bonus_f[n])))
        linha(M + 4, yy - RF, M + LW - 4)
        yy -= RF
    c.setFillColor(C["cinza"])
    c.rect(M, yy - 14, LW, 14, stroke=0, fill=1)
    c.setFillColor(C["azul"])
    c.setFont(B, 8)
    c.drawString(M + 6, yy - 9.5, "TOTAL — MÃO DE OBRA DA SEMANA" + (" (inclui bônus)" if st["bonus"] else ""))
    c.drawCentredString(M + 0.74 * LW, yy - 9.5, num(st["diarias"]) + (f" + {num(st['bonus_d'])}" if st["bonus"] else ""))
    c.drawRightString(M + LW - 6, yy - 9.5, brl(st["mo"]))

    # ---- esquerda, abaixo: acumulado da obra
    ac_y = ry - BLH - 8
    caixa(M, ac_y - ACH, LW, ACH)
    faixa(M, ac_y, LW, 13, C["azul_claro"], "ACUMULADO DA OBRA", 8, cor_txt=C["azul"])
    yy = ac_y - 13
    c.setFillColor(C["cinza"])
    c.rect(M, yy - 10, LW, 10, stroke=0, fill=1)
    c.setFillColor(C["azul"])
    c.setFont(B, 6.5)
    c.drawString(M + 6, yy - 7.2, "PERÍODO")
    c.drawCentredString(M + 0.42 * LW, yy - 7.2, "DIÁRIAS")
    c.drawCentredString(M + 0.58 * LW, yy - 7.2, "MÃO DE OBRA")
    c.drawCentredString(M + 0.76 * LW, yy - 7.2, "ALIM. / OUTROS")
    c.drawRightString(M + LW - 6, yy - 7.2, "TOTAL")
    yy -= 10
    acum = dict(diarias=0, mo=0, alim=0, total=0)
    for s_, sst in anteriores + [(sem, st)]:
        a, b = (ini, fim) if s_ is sem else (s_["_inicio"], s_["_fim"])
        c.setFillColor(C["texto"])
        c.setFont(R, F(7))
        c.drawString(M + 6, yy - RA + 2.5 * k, f"Semana {s_['numero']} — {a.strftime('%d/%m')} a {b.strftime('%d/%m')}")
        c.drawCentredString(M + 0.42 * LW, yy - RA + 2.5 * k, num(sst["diarias"]))
        c.drawCentredString(M + 0.58 * LW, yy - RA + 2.5 * k, brl(sst["mo"]))
        c.drawCentredString(M + 0.76 * LW, yy - RA + 2.5 * k, brl(sst["alim"]))
        c.setFont(B, F(7))
        c.drawRightString(M + LW - 6, yy - RA + 2.5 * k, brl(sst["total"]))
        linha(M + 4, yy - RA, M + LW - 4)
        yy -= RA
        for kk in acum:
            acum[kk] += sst[kk]
    c.setFillColor(C["verde"])
    c.rect(M, yy - 13, LW, 13, stroke=0, fill=1)
    c.setFillColor(C["azul"])
    c.setFont(B, 8)
    c.drawString(M + 6, yy - 9, "TOTAL ACUMULADO DA OBRA")
    c.drawCentredString(M + 0.42 * LW, yy - 9, num(acum["diarias"]))
    c.drawCentredString(M + 0.58 * LW, yy - 9, brl(acum["mo"]))
    c.drawCentredString(M + 0.76 * LW, yy - 9, brl(acum["alim"]))
    c.drawRightString(M + LW - 6, yy - 9, brl(acum["total"]))

    # ---- direita: quentinhas por valor (+ extras)
    BQ = 13 + 11 + 12 * len(faixas) + 14 + (12 * len(ext_sem) + 14 if ext_sem else 0)
    caixa(RX, ry - BQ, RW, BQ)
    faixa(RX, ry, RW, 13, C["azul_claro"],
          "QUENTINHAS — SEPARADO POR VALOR UNITÁRIO" + (" · E OUTROS" if ext_sem else ""), 8, cor_txt=C["azul"])
    yy = ry - 13
    c.setFillColor(C["cinza"])
    c.rect(RX, yy - 11, RW, 11, stroke=0, fill=1)
    c.setFillColor(C["azul"])
    c.setFont(B, 6.5)
    c.drawString(RX + 6, yy - 7.5, "VALOR UNITÁRIO")
    c.drawString(RX + 0.32 * RW, yy - 7.5, "DIAS")
    c.drawCentredString(RX + 0.76 * RW, yy - 7.5, "QTD.")
    c.drawRightString(RX + RW - 6, yy - 7.5, "CUSTO")
    yy -= 11

    def linha_q(lab, dias_txt, qtd_txt, custo):
        nonlocal yy
        c.setFillColor(C["texto"])
        c.setFont(B, 7.5)
        c.drawString(RX + 6, yy - 8, lab)
        c.setFont(R, 6.2)
        c.drawString(RX + 0.32 * RW, yy - 8, dias_txt)
        c.setFont(R, 7.5)
        c.drawCentredString(RX + 0.76 * RW, yy - 8, qtd_txt)
        c.setFont(B, 7.5)
        c.drawRightString(RX + RW - 6, yy - 8, brl(custo))
        linha(RX + 4, yy - 11, RX + RW - 4)
        yy -= 12

    def total_q(lab, qtd_txt, custo):
        nonlocal yy
        c.setFillColor(C["cinza"])
        c.rect(RX, yy - 14, RW, 14, stroke=0, fill=1)
        c.setFillColor(C["azul"])
        c.setFont(B, 7.5)
        c.drawString(RX + 6, yy - 9.5, lab)
        if qtd_txt:
            c.drawCentredString(RX + 0.76 * RW, yy - 9.5, qtd_txt)
        c.drawRightString(RX + RW - 6, yy - 9.5, brl(custo))
        yy -= 14

    for v, fx in faixas.items():
        linha_q(brl(v), ", ".join(fx["dias"]), f"{fx['qtd']} un.", fx["qtd"] * v)
    total_q("TOTAL DE QUENTINHAS DA SEMANA", f"{st['quent']} un.", st["quent_custo"])
    if ext_sem:
        for desc, g in ext_sem.items():
            linha_q(desc, ", ".join(g["dias"]), qtd_txt(g), g["custo"])
        total_q("TOTAL — ALIMENTAÇÃO E OUTROS", "", st["alim"])

    # ---- direita, abaixo: gasto geral
    yy = ry - BQ - 8
    c.setFillColor(C["texto"])
    c.setFont(R, 8)
    c.drawString(RX + 6, yy - 9, "Total de mão de obra")
    c.setFont(B, 8)
    c.drawRightString(RX + RW - 6, yy - 9, brl(st["mo"]))
    c.setFont(R, 8)
    c.drawString(RX + 6, yy - 20, "Total de alimentação" + (" e outros" if ext_sem else ""))
    c.setFont(B, 8)
    c.drawRightString(RX + RW - 6, yy - 20, brl(st["alim"]))
    c.setStrokeColor(C["linha"])
    c.setLineWidth(0.4)
    c.line(RX + 4, yy - 24, RX + RW - 4, yy - 24)
    c.setFillColor(C["verde"])
    c.rect(RX, yy - 47, RW, 19, stroke=0, fill=1)
    caixa(RX, yy - 47, RW, 19, lw=1.2)
    c.setFillColor(C["azul"])
    c.setFont(B, 11)
    c.drawString(RX + 8, yy - 41, "GASTO GERAL DA SEMANA")
    c.drawRightString(RX + RW - 8, yy - 41, brl(st["total"]))

    # ---- rodapé
    c.setFillColor(C["rodape"])
    c.setFont(O, 6.5)
    c.drawString(M, M - 12, f"{obra['empresa']} · Resp. téc. {obra['responsavel']} · "
                            f"Valores de diária e quentinha conforme informado pela obra · "
                            f"Emitido em {dt.date.today().strftime('%d/%m/%Y')}")
    c.drawRightString(W - M, M - 12, "Página 1 de 1")
    c.showPage()
    c.save()
    return nome_arq, st


# ===================================================================== XLSX
def gerar_xlsx(j, numero, saida, parcial=False):
    sem, dias, anteriores = semana_info(j, numero)
    equipe = j["_equipe"]
    eq_sem = equipe_da_semana(j, dias)
    obra = j["obra"]
    ini, fim = dias[0]["_data"], dias[-1]["_data"]
    A = FONTE_XLSX
    TIT = Font(name=A, size=13, bold=True, color="1F3864")
    SUB = Font(name=A, size=9, italic=True, color="595959")
    NEG = Font(name=A, size=10, bold=True)
    PRETO = Font(name=A, size=10)
    AZUL = Font(name=A, size=10, color="0000FF")
    VERDE = Font(name=A, size=10, color="008000")
    BRANCO = Font(name=A, size=10, bold=True, color="FFFFFF")
    BARRA = Font(name=A, size=10, bold=True, color="1F3864")
    PEQ = Font(name=A, size=9, color="404040")
    F_HDR = PatternFill("solid", fgColor="1F3864")
    F_BARRA = PatternFill("solid", fgColor="D9E1F2")
    F_IN = PatternFill("solid", fgColor=COR["input_xlsx"])
    F_SUB = PatternFill("solid", fgColor="EDEDED")
    F_TOT = PatternFill("solid", fgColor="C6E0B4")
    thin = Side(style="thin", color="BFBFBF")
    med = Side(style="medium", color="1F3864")
    BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
    BOXM = Border(left=med, right=med, top=med, bottom=med)
    MOEDA = 'R$ #,##0.00;-R$ #,##0.00;"-"'
    NUM = '0.0;-0.0;"-"'
    INT = '0;-0;"-"'
    DATA = "DD/MM/YYYY"
    OBRA = f"Obra: {obra['nome']}  ·  {obra['endereco']}  ·  {obra['empresa']}"
    N = len(eq_sem)
    L0, PAR0 = 6, 8
    LSUB = L0 + N

    def hdr(ws, row, titulos):
        for i, t in enumerate(titulos, start=1):
            cel = ws.cell(row=row, column=i, value=t)
            cel.font, cel.fill, cel.border = BRANCO, F_HDR, BOX
            cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 28

    def barra(ws, row, txt, span):
        ws[f"A{row}"] = txt
        ws[f"A{row}"].font, ws[f"A{row}"].fill = BARRA, F_BARRA
        ws.merge_cells(f"A{row}:{span}{row}")

    wb = Workbook()
    # ---------------- Parâmetros
    p = wb.active
    p.title = "Parâmetros"
    p["A1"] = f"{obra['empresa'].upper()} — CONTROLE DE DIÁRIAS E CUSTOS"
    p["A1"].font = TIT
    p["A2"] = OBRA
    p["A2"].font = SUB
    p["A4"] = "Semana de referência"
    p["A4"].font = NEG
    p["C4"], p["D4"], p["E4"] = ini, "a", fim
    for cc in ("C4", "E4"):
        p[cc].number_format, p[cc].font, p[cc].fill = DATA, AZUL, F_IN
    p["D4"].alignment = Alignment(horizontal="center")
    barra(p, 6, "EQUIPE E VALORES DE DIÁRIA", "D")
    hdr(p, 7, ["Funcionário", "Função", "Diária (R$)", "Situação"])
    for i, e in enumerate(eq_sem):
        r = PAR0 + i
        sit = "Ativo" if e["status"] == "ativo" else f"Desligado em {d(e['saida']).strftime('%d/%m')}"
        p[f"A{r}"], p[f"B{r}"], p[f"C{r}"], p[f"D{r}"] = e["nome"], e["funcao"], e["diaria"], sit
        p[f"A{r}"].font = p[f"B{r}"].font = PRETO
        p[f"C{r}"].font, p[f"C{r}"].fill, p[f"C{r}"].number_format = AZUL, F_IN, MOEDA
        p[f"D{r}"].font = PEQ
        for cc in "ABCD":
            p[f"{cc}{r}"].border = BOX
    ra = PAR0 + N + 1
    barra(p, ra, "ALIMENTAÇÃO — HISTÓRICO DE PREÇO DA QUENTINHA", "D")
    hdr(p, ra + 1, ["Vigência", "", "Valor unitário (R$)", "Observação"])
    hist = historico_precos(j, fim)
    for i, h in enumerate(hist):
        r = ra + 2 + i
        vig = (f"Dia {h['inicio'].strftime('%d/%m')}" if h["inicio"] == h["fim"]
               else f"De {h['inicio'].strftime('%d/%m')} a {h['fim'].strftime('%d/%m')}")
        p[f"A{r}"] = vig
        p[f"A{r}"].font = PRETO
        p.merge_cells(f"A{r}:B{r}")
        p[f"C{r}"] = h["valor"]
        p[f"C{r}"].font, p[f"C{r}"].fill, p[f"C{r}"].number_format = AZUL, F_IN, MOEDA
        p[f"D{r}"] = "Valor corrente" if i == len(hist) - 1 else ""
        p[f"D{r}"].font = PEQ
        for cc in "ABCD":
            p[f"{cc}{r}"].border = BOX
    rh = ra + 2 + len(hist) + 1
    barra(p, rh, "COMO USAR", "D")
    for i, t in enumerate([
        "Células amarelas com texto azul são de preenchimento. As demais são fórmulas — não digite por cima.",
        "Nas abas de dia, coluna 'Presença': 1 = diária inteira · 0,5 = meia diária · vazio = não trabalhou.",
        "O valor da quentinha é lançado DENTRO de cada aba de dia — cada data guarda o preço que valeu naquele dia.",
        "A quantidade de quentinhas é livre: pode ser maior que o número de presentes, quando houver unidade adicional.",
        "Outros consumos do dia (gelo, entrega de quentinha) entram como linhas próprias na seção 2 de cada aba de dia.",
        "A aba 'Resumo Semanal' consolida os dias, o gasto por funcionário, as quentinhas por valor unitário e o acumulado da obra.",
        "Aparecem aqui apenas os funcionários que tiveram presença nesta semana.",
    ]):
        p[f"A{rh+1+i}"] = t
        p[f"A{rh+1+i}"].font = PEQ
    for cc, w in {"A": 46, "B": 16, "C": 17, "D": 40, "E": 14}.items():
        p.column_dimensions[cc].width = w
    p.sheet_view.showGridLines = False

    # ---------------- abas de dia
    REFS = {}
    for dia in dias:
        data = dia["_data"]
        aba = f"{DIAS_CURTO[data.weekday()]} {data.strftime('%d-%m')}"
        s = wb.create_sheet(aba)
        s["A1"] = f"{DIAS_SEMANA[data.weekday()]} — {data.strftime('%d/%m/%Y')}"
        s["A1"].font = TIT
        s.merge_cells("A1:E1")
        s["A2"] = OBRA
        s["A2"].font = SUB
        s.merge_cells("A2:E2")
        barra(s, 4, "1. MÃO DE OBRA — DETALHAMENTO POR FUNCIONÁRIO", "E")
        hdr(s, 5, ["Funcionário", "Função", "Diária (R$)", "Presença", "Valor do Dia (R$)"])
        s.row_dimensions[5].height = 26
        for i, e in enumerate(eq_sem):
            r = L0 + i
            s[f"A{r}"], s[f"B{r}"] = e["nome"], e["funcao"]
            s[f"A{r}"].font = s[f"B{r}"].font = PRETO
            s[f"C{r}"] = f"='Parâmetros'!$C${PAR0+i}"
            s[f"C{r}"].font, s[f"C{r}"].number_format = VERDE, MOEDA
            if e["nome"] in dia["presencas"]:
                s[f"D{r}"] = dia["presencas"][e["nome"]]
            s[f"D{r}"].font, s[f"D{r}"].fill, s[f"D{r}"].number_format = AZUL, F_IN, NUM
            s[f"D{r}"].alignment = Alignment(horizontal="center")
            s[f"E{r}"] = f"=C{r}*D{r}"
            s[f"E{r}"].number_format, s[f"E{r}"].font = MOEDA, PRETO
            for cc in "ABCDE":
                s[f"{cc}{r}"].border = BOX
        bon = bonus_de(dia, equipe)
        for kb, (n, q, v) in enumerate(bon):
            rb_ = LSUB + kb
            i_e = [e["nome"] for e in eq_sem].index(n)
            s[f"A{rb_}"] = f"Bônus de produção — {n}"
            s[f"A{rb_}"].font = PRETO
            s[f"B{rb_}"] = "bônus"
            s[f"B{rb_}"].font = PEQ
            s[f"C{rb_}"] = f"='Parâmetros'!$C${PAR0+i_e}"
            s[f"C{rb_}"].font, s[f"C{rb_}"].number_format = VERDE, MOEDA
            s[f"D{rb_}"] = q
            s[f"D{rb_}"].font, s[f"D{rb_}"].fill, s[f"D{rb_}"].number_format = AZUL, F_IN, NUM
            s[f"D{rb_}"].alignment = Alignment(horizontal="center")
            s[f"E{rb_}"] = f"=C{rb_}*D{rb_}"
            s[f"E{rb_}"].number_format, s[f"E{rb_}"].font = MOEDA, PRETO
            for cc in "ABCDE":
                s[f"{cc}{rb_}"].border = BOX
        LSUBd = LSUB + len(bon)
        s[f"A{LSUBd}"] = "SUBTOTAL — MÃO DE OBRA" + (" (inclui bônus)" if bon else "")
        s[f"A{LSUBd}"].font = NEG
        s.merge_cells(f"A{LSUBd}:B{LSUBd}")
        s[f"C{LSUBd}"] = "Diárias:"
        s[f"C{LSUBd}"].font = NEG
        s[f"C{LSUBd}"].alignment = Alignment(horizontal="right")
        s[f"D{LSUBd}"] = f"=SUM(D{L0}:D{LSUB-1})"
        s[f"D{LSUBd}"].font, s[f"D{LSUBd}"].number_format = NEG, NUM
        s[f"D{LSUBd}"].alignment = Alignment(horizontal="center")
        s[f"E{LSUBd}"] = f"=SUM(E{L0}:E{LSUBd-1})"
        s[f"E{LSUBd}"].font, s[f"E{LSUBd}"].number_format = NEG, MOEDA
        for cc in "ABCDE":
            s[f"{cc}{LSUBd}"].fill, s[f"{cc}{LSUBd}"].border = F_SUB, BOX
        rq = LSUBd + 2
        barra(s, rq, "2. ALIMENTAÇÃO E OUTROS CONSUMOS DO DIA", "E")
        hdr(s, rq + 1, ["Descrição", "", "Valor unitário (R$)", "Quantidade", "Valor Total (R$)"])
        s.row_dimensions[rq + 1].height = 26
        rq2 = rq + 2
        itens = [("Quentinhas fornecidas", dia["valor_quentinha"], dia["quentinhas"])]
        itens += [(desc, v, q) for desc, q, v in extras_de(dia)]
        for kx, (desc, v, q) in enumerate(itens):
            rk = rq2 + kx
            s[f"A{rk}"] = desc
            s[f"A{rk}"].font = PRETO
            s.merge_cells(f"A{rk}:B{rk}")
            s[f"C{rk}"] = v
            s[f"C{rk}"].font, s[f"C{rk}"].fill, s[f"C{rk}"].number_format = AZUL, F_IN, MOEDA
            s[f"D{rk}"] = q
            s[f"D{rk}"].font, s[f"D{rk}"].fill, s[f"D{rk}"].number_format = AZUL, F_IN, INT
            s[f"D{rk}"].alignment = Alignment(horizontal="center")
            s[f"E{rk}"] = f"=C{rk}*D{rk}"
            s[f"E{rk}"].number_format, s[f"E{rk}"].font = MOEDA, PRETO
            for cc in "ABCDE":
                s[f"{cc}{rk}"].border = BOX
        rqf = rq2 + len(itens) - 1
        if len(itens) > 1:
            rqf += 1
            s[f"A{rqf}"] = "SUBTOTAL — ALIMENTAÇÃO E OUTROS"
            s[f"A{rqf}"].font = NEG
            s.merge_cells(f"A{rqf}:D{rqf}")
            s[f"E{rqf}"] = f"=SUM(E{rq2}:E{rqf-1})"
            s[f"E{rqf}"].number_format, s[f"E{rqf}"].font = MOEDA, NEG
            for cc in "ABCDE":
                s[f"{cc}{rqf}"].fill, s[f"{cc}{rqf}"].border = F_SUB, BOX
        rr = rqf + 2
        barra(s, rr, f"3. RESUMO DO DIA — {DIAS_SEMANA[data.weekday()]} {data.strftime('%d/%m')}", "E")
        for i, (lab, fml, fmt) in enumerate([
                ("Total de diárias trabalhadas", f"=D{LSUBd}", NUM),
                ("Quentinhas fornecidas (un.)", f"=D{rq2}", INT),
                ("Custo de mão de obra" + (" (inclui bônus)" if bon else ""), f"=E{LSUBd}", MOEDA),
                ("Custo de alimentação" + (" e outros" if len(itens) > 1 else ""), f"=E{rqf}", MOEDA)]):
            r = rr + 1 + i
            s[f"A{r}"] = lab
            s[f"A{r}"].font = PRETO
            s.merge_cells(f"A{r}:D{r}")
            s[f"E{r}"] = fml
            s[f"E{r}"].number_format, s[f"E{r}"].font = fmt, PRETO
            for cc in "ABCDE":
                s[f"{cc}{r}"].border = BOX
        rt = rr + 5
        s[f"A{rt}"] = "TOTAL DO DIA"
        s[f"A{rt}"].font = Font(name=A, size=11, bold=True)
        s.merge_cells(f"A{rt}:D{rt}")
        s[f"E{rt}"] = f"=E{rr+3}+E{rr+4}"
        s[f"E{rt}"].font, s[f"E{rt}"].number_format = Font(name=A, size=11, bold=True), MOEDA
        for cc in "ABCDE":
            s[f"{cc}{rt}"].fill, s[f"{cc}{rt}"].border = F_TOT, BOXM
        if dia.get("observacao"):
            s[f"A{rt+2}"] = f"Observação: {dia['observacao']}"
            s[f"A{rt+2}"].font = Font(name=A, size=9, italic=True, color="C00000")
            s.merge_cells(f"A{rt+2}:E{rt+3}")
            s[f"A{rt+2}"].alignment = Alignment(wrap_text=True, vertical="top")
        for cc, w in {"A": 24, "B": 15, "C": 17, "D": 12, "E": 19}.items():
            s.column_dimensions[cc].width = w
        s.sheet_view.showGridLines = False
        s.page_setup.orientation = "portrait"
        s.page_setup.fitToWidth = s.page_setup.fitToHeight = 1
        s.sheet_properties.pageSetUpPr.fitToPage = True
        REFS[aba] = dict(diarias=f"$E${rr+1}", quent=f"$E${rr+2}", mo=f"$E${rr+3}",
                         alim=f"$E${rr+4}", tot=f"$E${rt}", qq=f"$D${rq2}", qcusto=f"$E${rq2}",
                         extras={desc: f"$E${rq2+1+kx}" for kx, (desc, _, _) in enumerate(itens[1:])},
                         bonus={n: f"$E${LSUB+kb}" for kb, (n, _, _) in enumerate(bon)})

    # ---------------- Resumo Semanal
    rs = wb.create_sheet("Resumo Semanal")
    rs["A1"] = (f"RESUMO GERAL DA SEMANA — {ini.strftime('%d/%m')} A {fim.strftime('%d/%m/%Y')}"
                + ("  (PARCIAL)" if parcial else ""))
    rs["A1"].font = Font(name=A, size=14, bold=True, color="1F3864")
    rs.merge_cells("A1:G1")
    rs["A2"] = OBRA
    rs["A2"].font = SUB
    rs.merge_cells("A2:G2")
    barra(rs, 4, "1. FECHAMENTO POR DIA", "G")
    hdr(rs, 5, ["Dia", "Data", "Diárias", "Mão de Obra (R$)", "Quentinhas (un.)", "Alim. e outros (R$)", "Total do Dia (R$)"])
    abas = list(REFS)
    for i, (aba, dia) in enumerate(zip(abas, dias)):
        r = 6 + i
        rs[f"A{r}"] = DIAS_CURTO[dia["_data"].weekday()]
        rs[f"B{r}"] = dia["_data"]
        rs[f"B{r}"].number_format = DATA
        for col, key, fmt in [("C", "diarias", NUM), ("D", "mo", MOEDA), ("E", "quent", INT),
                              ("F", "alim", MOEDA), ("G", "tot", MOEDA)]:
            rs[f"{col}{r}"] = f"='{aba}'!{REFS[aba][key]}"
            rs[f"{col}{r}"].number_format = fmt
            rs[f"{col}{r}"].font = NEG if col == "G" else PRETO
        rs[f"A{r}"].font = rs[f"B{r}"].font = PRETO
        for cc in "ABCDEFG":
            rs[f"{cc}{r}"].border = BOX
    rtd = 6 + len(dias)
    rs[f"A{rtd}"] = "TOTAL DA SEMANA"
    rs[f"A{rtd}"].font = NEG
    rs.merge_cells(f"A{rtd}:B{rtd}")
    for col, fmt in [("C", NUM), ("D", MOEDA), ("E", INT), ("F", MOEDA), ("G", MOEDA)]:
        rs[f"{col}{rtd}"] = f"=SUM({col}6:{col}{rtd-1})"
        rs[f"{col}{rtd}"].number_format, rs[f"{col}{rtd}"].font = fmt, NEG
    for cc in "ABCDEFG":
        rs[f"{cc}{rtd}"].fill, rs[f"{cc}{rtd}"].border = F_SUB, BOX

    rb = rtd + 2
    barra(rs, rb, "2. GASTOS POR FUNCIONÁRIO", "G")
    hdr(rs, rb + 1, ["Funcionário", "Função", "Diária (R$)", "Diárias na Semana", "", "", "Total a Pagar (R$)"])
    rs.merge_cells(f"D{rb+1}:F{rb+1}")
    RF0 = rb + 2
    for i, e in enumerate(eq_sem):
        r = RF0 + i
        rs[f"A{r}"], rs[f"B{r}"] = e["nome"], e["funcao"]
        rs[f"A{r}"].font = rs[f"B{r}"].font = PRETO
        rs[f"C{r}"] = f"='Parâmetros'!$C${PAR0+i}"
        rs[f"C{r}"].font, rs[f"C{r}"].number_format = VERDE, MOEDA
        rs[f"D{r}"] = "=" + "+".join([f"'{aba}'!$D${L0+i}" for aba in abas])
        rs[f"D{r}"].number_format, rs[f"D{r}"].font = NUM, PRETO
        rs[f"D{r}"].alignment = Alignment(horizontal="center")
        rs.merge_cells(f"D{r}:F{r}")
        refs_b = [f"'{aba}'!{REFS[aba]['bonus'][e['nome']]}" for aba in abas if e["nome"] in REFS[aba]["bonus"]]
        rs[f"G{r}"] = f"=C{r}*D{r}" + ("+" + "+".join(refs_b) if refs_b else "")
        rs[f"G{r}"].number_format, rs[f"G{r}"].font = MOEDA, NEG
        if refs_b:
            rs[f"B{r}"] = e["funcao"] + " · c/ bônus"
        for cc in "ABCDEFG":
            rs[f"{cc}{r}"].border = BOX
    rfm = RF0 + N
    rs[f"A{rfm}"] = "TOTAL — MÃO DE OBRA DA SEMANA"
    rs[f"A{rfm}"].font = NEG
    rs.merge_cells(f"A{rfm}:C{rfm}")
    rs[f"D{rfm}"] = f"=SUM(D{RF0}:D{rfm-1})"
    rs[f"D{rfm}"].number_format, rs[f"D{rfm}"].font = NUM, NEG
    rs[f"D{rfm}"].alignment = Alignment(horizontal="center")
    rs.merge_cells(f"D{rfm}:F{rfm}")
    rs[f"G{rfm}"] = f"=SUM(G{RF0}:G{rfm-1})"
    rs[f"G{rfm}"].number_format, rs[f"G{rfm}"].font = MOEDA, NEG
    for cc in "ABCDEFG":
        rs[f"{cc}{rfm}"].fill, rs[f"{cc}{rfm}"].border = F_SUB, BOX

    faixas = faixas_quentinha(dias)
    ext_sem = extras_agrupados(dias)
    rq = rfm + 2
    barra(rs, rq, "3. GASTOS COM QUENTINHAS — SEPARADO POR VALOR UNITÁRIO (E OUTROS CONSUMOS)", "G")
    hdr(rs, rq + 1, ["Valor unitário / item", "Dias em que foi usado", "", "", "Quantidade (un.)", "", "Custo (R$)"])
    rs.merge_cells(f"B{rq+1}:D{rq+1}")
    rs.merge_cells(f"E{rq+1}:F{rq+1}")
    RQ0 = rq + 2
    for i, (v, fx) in enumerate(faixas.items()):
        r = RQ0 + i
        dd = [(aba, dia) for aba, dia in zip(abas, dias) if dia["valor_quentinha"] == v and dia["quentinhas"]]
        rs[f"A{r}"] = v
        rs[f"A{r}"].number_format, rs[f"A{r}"].font = MOEDA, NEG
        rs[f"B{r}"] = ", ".join(fx["dias"])
        rs[f"B{r}"].font = PRETO
        rs.merge_cells(f"B{r}:D{r}")
        rs[f"E{r}"] = "=" + "+".join([f"'{aba}'!{REFS[aba]['qq']}" for aba, _ in dd])
        rs[f"E{r}"].number_format, rs[f"E{r}"].font = INT, PRETO
        rs[f"E{r}"].alignment = Alignment(horizontal="center")
        rs.merge_cells(f"E{r}:F{r}")
        rs[f"G{r}"] = "=" + "+".join([f"'{aba}'!{REFS[aba]['qcusto']}" for aba, _ in dd])
        rs[f"G{r}"].number_format, rs[f"G{r}"].font = MOEDA, NEG
        for cc in "ABCDEFG":
            rs[f"{cc}{r}"].border = BOX
    for i, (desc, g) in enumerate(ext_sem.items()):
        r = RQ0 + len(faixas) + i
        rs[f"A{r}"] = desc
        rs[f"A{r}"].font = NEG
        rs[f"B{r}"] = ", ".join(g["dias"]) + f"  ({qtd_txt(g)})"
        rs[f"B{r}"].font = PRETO
        rs.merge_cells(f"B{r}:D{r}")
        rs[f"E{r}"] = g["qtd"]
        rs[f"E{r}"].number_format, rs[f"E{r}"].font = INT, PRETO
        rs[f"E{r}"].alignment = Alignment(horizontal="center")
        rs.merge_cells(f"E{r}:F{r}")
        refs_e = [f"'{aba}'!{REFS[aba]['extras'][desc]}" for aba in abas if desc in REFS[aba]["extras"]]
        rs[f"G{r}"] = "=" + "+".join(refs_e)
        rs[f"G{r}"].number_format, rs[f"G{r}"].font = MOEDA, NEG
        for cc in "ABCDEFG":
            rs[f"{cc}{r}"].border = BOX
    rqt = RQ0 + len(faixas) + len(ext_sem)
    rs[f"A{rqt}"] = "TOTAL — ALIMENTAÇÃO E OUTROS" if ext_sem else "TOTAL DE QUENTINHAS DA SEMANA"
    rs[f"A{rqt}"].font = NEG
    rs.merge_cells(f"A{rqt}:D{rqt}")
    rs[f"E{rqt}"] = f"=SUM(E{RQ0}:E{RQ0+len(faixas)-1})"
    rs[f"E{rqt}"].number_format, rs[f"E{rqt}"].font = INT, NEG
    rs[f"E{rqt}"].alignment = Alignment(horizontal="center")
    rs.merge_cells(f"E{rqt}:F{rqt}")
    rs[f"G{rqt}"] = f"=SUM(G{RQ0}:G{rqt-1})"
    rs[f"G{rqt}"].number_format, rs[f"G{rqt}"].font = MOEDA, NEG
    for cc in "ABCDEFG":
        rs[f"{cc}{rqt}"].fill, rs[f"{cc}{rqt}"].border = F_SUB, BOX

    rg = rqt + 2
    barra(rs, rg, "4. GASTO GERAL DA SEMANA", "G")
    for i, (lab, fml) in enumerate([("Total de mão de obra", f"=G{rfm}"),
                                    ("Total de alimentação e outros", f"=G{rqt}")]):
        r = rg + 1 + i
        rs[f"A{r}"] = lab
        rs[f"A{r}"].font = PRETO
        rs.merge_cells(f"A{r}:F{r}")
        rs[f"G{r}"] = fml
        rs[f"G{r}"].number_format, rs[f"G{r}"].font = MOEDA, PRETO
        for cc in "ABCDEFG":
            rs[f"{cc}{r}"].border = BOX
    rgt = rg + 3
    rs[f"A{rgt}"] = "GASTO GERAL DA SEMANA"
    rs[f"A{rgt}"].font = Font(name=A, size=12, bold=True)
    rs.merge_cells(f"A{rgt}:F{rgt}")
    rs[f"G{rgt}"] = f"=G{rg+1}+G{rg+2}"
    rs[f"G{rgt}"].font, rs[f"G{rgt}"].number_format = Font(name=A, size=12, bold=True), MOEDA
    for cc in "ABCDEFG":
        rs[f"{cc}{rgt}"].fill, rs[f"{cc}{rgt}"].border = F_TOT, BOXM

    rac = rgt + 2
    barra(rs, rac, "5. ACUMULADO DA OBRA", "G")
    hdr(rs, rac + 1, ["Período", "", "Diárias", "Mão de Obra (R$)", "Quentinhas (un.)", "Alim. e outros (R$)", "Total (R$)"])
    rs.merge_cells(f"A{rac+1}:B{rac+1}")
    RA0 = rac + 2
    for i, (s_, sst) in enumerate(anteriores):
        r = RA0 + i
        rs[f"A{r}"] = f"Semana {s_['numero']} — {s_['_inicio'].strftime('%d/%m')} a {s_['_fim'].strftime('%d/%m')}"
        rs[f"A{r}"].font = PRETO
        rs.merge_cells(f"A{r}:B{r}")
        for col, val, fmt in [("C", sst["diarias"], NUM), ("D", sst["mo"], MOEDA), ("E", sst["quent"], INT),
                              ("F", sst["alim"], MOEDA), ("G", sst["total"], MOEDA)]:
            rs[f"{col}{r}"] = val
            rs[f"{col}{r}"].number_format, rs[f"{col}{r}"].font, rs[f"{col}{r}"].fill = fmt, AZUL, F_IN
    r3 = RA0 + len(anteriores)
    rs[f"A{r3}"] = f"Semana {numero} — {ini.strftime('%d/%m')} a {fim.strftime('%d/%m')}"
    rs[f"A{r3}"].font = PRETO
    rs.merge_cells(f"A{r3}:B{r3}")
    for col, fmt in [("C", NUM), ("D", MOEDA), ("E", INT), ("F", MOEDA), ("G", MOEDA)]:
        rs[f"{col}{r3}"] = f"={col}{rtd}"
        rs[f"{col}{r3}"].number_format, rs[f"{col}{r3}"].font = fmt, PRETO
    r4 = r3 + 1
    rs[f"A{r4}"] = "TOTAL ACUMULADO DA OBRA"
    rs[f"A{r4}"].font = NEG
    rs.merge_cells(f"A{r4}:B{r4}")
    for col, fmt in [("C", NUM), ("D", MOEDA), ("E", INT), ("F", MOEDA), ("G", MOEDA)]:
        rs[f"{col}{r4}"] = f"=SUM({col}{RA0}:{col}{r3})"
        rs[f"{col}{r4}"].number_format, rs[f"{col}{r4}"].font = fmt, NEG
    for r in range(RA0, r4 + 1):
        for cc in "ABCDEFG":
            rs[f"{cc}{r}"].border = BOX
    for cc in "ABCDEFG":
        rs[f"{cc}{r4}"].fill = F_TOT

    notas = "; ".join(f"{x['_data'].strftime('%d/%m')} — {x['observacao']}" for x in dias if x.get("observacao"))
    rn = r4 + 2
    rs[f"A{rn}"] = (f"Notas da semana: {notas}. Constam apenas os funcionários com presença na semana. "
                    "Os totais de mão de obra dos blocos 1 e 2 devem ser idênticos.")
    rs[f"A{rn}"].font = Font(name=A, size=9, italic=True, color="404040")
    rs.merge_cells(f"A{rn}:G{rn+2}")
    rs[f"A{rn}"].alignment = Alignment(wrap_text=True, vertical="top")
    for cc, w in {"A": 32, "B": 14, "C": 13, "D": 18, "E": 16, "F": 18, "G": 20}.items():
        rs.column_dimensions[cc].width = w
    rs.sheet_view.showGridLines = False
    rs.page_setup.orientation = "portrait"
    rs.page_setup.fitToWidth = rs.page_setup.fitToHeight = 1
    rs.sheet_properties.pageSetUpPr.fitToPage = True

    nome_arq = (f"Controle_Diarias_{obra['nome'].split('–')[0].strip().replace(' ', '_')}_"
                f"Semana{numero}_{ini.strftime('%d-%m')}-a-{fim.strftime('%d-%m')}.xlsx")
    wb.save(os.path.join(saida, nome_arq))
    return nome_arq


# ===================================================================== MAIN
if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--dados", required=True)
    ap.add_argument("--semana", type=int, required=True)
    ap.add_argument("--saida", default=".")
    ap.add_argument("--parcial", action="store_true", help="semana ainda em andamento")
    a = ap.parse_args()
    os.makedirs(a.saida, exist_ok=True)
    j = carregar(a.dados)
    pdf, st = gerar_pdf(j, a.semana, a.saida, a.parcial)
    xlsx = gerar_xlsx(j, a.semana, a.saida, a.parcial)
    print(f"PDF : {pdf}")
    print(f"XLSX: {xlsx}")
    print(f"Semana {a.semana}: {num(st['diarias'])} diárias · M.O. {brl(st['mo'])} · "
          f"{st['quent']} quentinhas · alimentação e outros {brl(st['alim'])} · TOTAL {brl(st['total'])}")
