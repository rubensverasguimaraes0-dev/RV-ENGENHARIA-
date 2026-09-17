# -*- coding: utf-8 -*-
"""
FECHAMENTO MENSAL — gerador de referência (RV Engenharia)

Uso:
    python relatorio_mensal.py --dados dados/obra_selecta.json --saida ./saida [--hoje AAAA-MM-DD]

Produz:
    Fechamento_Mensal_<obra>.pdf     (1 página, A4 paisagem — um card por mês)
    Fechamento_Mensal_<obra>.xlsx    (aba Resumo Mensal + uma aba por mês)

Regras: cada dia conta no seu mês do calendário. O mês corrente sai como PARCIAL.
Só aparecem funcionários com presença em pelo menos um mês.
"Alimentação e outros" = quentinhas + extras (gelo, entrega etc.).
"""
import argparse
import datetime as dt
import json
import os
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.colors import HexColor, white

D = dt.date
ap = argparse.ArgumentParser()
ap.add_argument("--dados", required=True)
ap.add_argument("--saida", default=".")
ap.add_argument("--hoje", default=None)
ARGS = ap.parse_args()
os.makedirs(ARGS.saida, exist_ok=True)
HOJE = D.fromisoformat(ARGS.hoje) if ARGS.hoje else D.today()
with open(ARGS.dados, encoding="utf-8") as _f:
    J = json.load(_f)
OBRA_J = J["obra"]
SLUG = OBRA_J["nome"].split("–")[0].strip().replace(" ", "_")

FUNCIONARIOS = [(e["nome"], e["funcao"], float(e["diaria"])) for e in J["equipe"]]
DIARIA = {f[0]: f[2] for f in FUNCIONARIOS}
FUNCAO = {f[0]: f[1] for f in FUNCIONARIOS}
ORDEM = [f[0] for f in FUNCIONARIOS]
SEMANAS = [(s["numero"], D.fromisoformat(s["inicio"]), D.fromisoformat(s["fim"])) for s in J["semanas"]]
# (data, presencas, qtd quentinhas, valor quentinha, extras[(desc, qtd, valor)])
L = sorted([(D.fromisoformat(x["data"]), x["presencas"], x["quentinhas"], float(x["valor_quentinha"]),
             [(e["descricao"], e.get("qtd", 1), float(e["valor"])) for e in x.get("extras", [])])
            for x in J["lancamentos"]], key=lambda x: x[0])

NOMES_MES = ["", "JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO", "JULHO", "AGOSTO",
             "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO"]
MESES = OrderedDict()
for (dd, *_r) in L:
    k = (dd.year, dd.month)
    if k in MESES:
        continue
    dias_mes = [x[0] for x in L if (x[0].year, x[0].month) == k]
    parcial = (k == (HOJE.year, HOJE.month))
    periodo = f"{min(dias_mes).strftime('%d/%m')} a {max(dias_mes).strftime('%d/%m')}" + (" — parcial" if parcial else "")
    MESES[k] = (f"{NOMES_MES[dd.month]}/{dd.year}", NOMES_MES[dd.month].capitalize(), periodo, parcial)


def brl(v):
    s = f"{v:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"R$ {s}"


def fmtn(v):
    return f"{v:g}".replace(".", ",")


def semana_de(d):
    for n, ini, fim in SEMANAS:
        if ini <= d <= fim:
            return n


def stats(dias):
    di = sum(sum(x[1].values()) for x in dias)
    mo = sum(sum(DIARIA[k] * v for k, v in x[1].items()) for x in dias)
    qq = sum(x[2] for x in dias)
    qc = sum(x[2] * x[3] for x in dias)
    ex = sum(q * v for x in dias for _, q, v in x[4])
    return dict(dias=len(dias), diarias=di, mo=mo, quent=qq, quent_custo=qc, extras=ex, alim=qc + ex, total=mo + qc + ex)


POR_MES = OrderedDict((k, [x for x in L if (x[0].year, x[0].month) == k]) for k in MESES)
STATS = {k: stats(v) for k, v in POR_MES.items()}
TOTAL = stats(L)

FUNC_MES = {n: {k: 0.0 for k in MESES} for n in ORDEM}
for k, dias in POR_MES.items():
    for _, pres, _, _, _ in dias:
        for n, f in pres.items():
            FUNC_MES[n][k] += f
FUNC_ATIVOS = [n for n in ORDEM if sum(FUNC_MES[n].values()) > 0]

VALORES = sorted({x[3] for x in L if x[2] > 0})
Q_MES = {v: {k: 0 for k in MESES} for v in VALORES}
for k, dias in POR_MES.items():
    for _, _, qq, vq, _ in dias:
        if qq:
            Q_MES[vq][k] += qq

EXTRAS_MES = OrderedDict()   # descricao -> {mes: (qtd, custo)}
for (dd, _p, _q, _v, ex) in L:
    k = (dd.year, dd.month)
    for desc, q, v in ex:
        EXTRAS_MES.setdefault(desc, {})
        a, b = EXTRAS_MES[desc].get(k, (0, 0.0))
        EXTRAS_MES[desc][k] = (a + q, b + q * v)

SEM_MES = OrderedDict()
for k, dias in POR_MES.items():
    grupos = OrderedDict()
    for x in dias:
        grupos.setdefault(semana_de(x[0]), []).append(x)
    SEM_MES[k] = [(n, g[0][0], g[-1][0], stats(g)) for n, g in grupos.items()]

# ================================================================== XLSX
A = "Arial"
TIT = Font(name=A, size=13, bold=True, color="1F3864")
SUB = Font(name=A, size=9, italic=True, color="595959")
NEG = Font(name=A, size=10, bold=True)
PRETO = Font(name=A, size=10)
BRANCO = Font(name=A, size=10, bold=True, color="FFFFFF")
BARRA = Font(name=A, size=10, bold=True, color="1F3864")
PEQ = Font(name=A, size=9, color="404040")
F_HDR = PatternFill("solid", fgColor="1F3864")
F_BARRA = PatternFill("solid", fgColor="D9E1F2")
F_SUB = PatternFill("solid", fgColor="EDEDED")
F_TOT = PatternFill("solid", fgColor="C6E0B4")
thin = Side(style="thin", color="BFBFBF")
med = Side(style="medium", color="1F3864")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
BOXM = Border(left=med, right=med, top=med, bottom=med)
MOEDA = 'R$ #,##0.00;-R$ #,##0.00;"-"'
NUM = '0.0;-0.0;"-"'
INT = '0;-0;"-"'
DATA = "DD/MM/YYYY"
OBRA = f"Obra: {OBRA_J['nome']}  ·  {OBRA_J['endereco']}  ·  {OBRA_J['empresa']}"
DIAS_SEM = ["Segunda", "Terça", "Quarta", "Quinta", "Sexta", "Sábado", "Domingo"]


def hdr(ws, row, titulos):
    for i, t in enumerate(titulos, start=1):
        cel = ws.cell(row=row, column=i, value=t)
        cel.font, cel.fill, cel.border = BRANCO, F_HDR, BOX
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 28


def barra(ws, row, txt, span):
    ws[f"A{row}"] = txt
    ws[f"A{row}"].font, ws[f"A{row}"].fill = BARRA, F_BARRA
    ws.merge_cells(f"A{row}:{span}{row}")


def celf(ws, ref, val, fmt=None, font=PRETO, fill=None, center=False, border=True):
    ws[ref] = val
    if fmt:
        ws[ref].number_format = fmt
    ws[ref].font = font
    if fill:
        ws[ref].fill = fill
    if center:
        ws[ref].alignment = Alignment(horizontal="center")
    if border:
        ws[ref].border = BOX


wb = Workbook()
wb.remove(wb.active)

for key, (nome_up, nome, periodo, parcial) in MESES.items():
    ws = wb.create_sheet(nome)
    ws["A1"] = f"FECHAMENTO MENSAL — {nome_up}" + ("  (PARCIAL)" if parcial else "")
    ws["A1"].font = TIT
    ws.merge_cells("A1:I1")
    ws["A2"] = f"{OBRA}  ·  Período: {periodo}"
    ws["A2"].font = SUB
    ws.merge_cells("A2:I2")

    barra(ws, 4, "1. LANÇAMENTOS DO MÊS", "I")
    hdr(ws, 5, ["Data", "Dia", "Semana", "Diárias", "Mão de Obra (R$)", "Quentinhas (un.)",
                "Valor unit. (R$)", "Outros (R$)", "Total do Dia (R$)"])
    r = 6
    for data, pres, qq, vq, ex in POR_MES[key]:
        mo = sum(DIARIA[k] * v for k, v in pres.items())
        celf(ws, f"A{r}", data, DATA)
        celf(ws, f"B{r}", DIAS_SEM[data.weekday()], center=True)
        celf(ws, f"C{r}", semana_de(data), center=True)
        celf(ws, f"D{r}", sum(pres.values()), NUM, center=True)
        celf(ws, f"E{r}", mo, MOEDA)
        celf(ws, f"F{r}", qq, INT, center=True)
        celf(ws, f"G{r}", vq, MOEDA)
        celf(ws, f"H{r}", sum(q * v for _, q, v in ex), MOEDA)
        celf(ws, f"I{r}", f"=E{r}+F{r}*G{r}+H{r}", MOEDA, NEG)
        r += 1
    rt = r
    ws[f"A{rt}"] = f"TOTAL DE {nome_up}"
    ws[f"A{rt}"].font = NEG
    ws.merge_cells(f"A{rt}:C{rt}")
    for col, fmt in [("D", NUM), ("E", MOEDA), ("F", INT), ("H", MOEDA), ("I", MOEDA)]:
        celf(ws, f"{col}{rt}", f"=SUM({col}6:{col}{rt-1})", fmt, NEG)
    celf(ws, f"G{rt}", f"=I{rt}-E{rt}-H{rt}", MOEDA, NEG)
    for cc in "ABCDEFGHI":
        ws[f"{cc}{rt}"].fill, ws[f"{cc}{rt}"].border = F_TOT, BOX
    ws[f"A{rt+1}"] = "Na linha de total, a coluna 'Valor unit.' mostra o custo total das quentinhas do mês."
    ws[f"A{rt+1}"].font = PEQ
    ws.merge_cells(f"A{rt+1}:I{rt+1}")

    r = rt + 3
    barra(ws, r, "2. GASTOS POR FUNCIONÁRIO NO MÊS", "H")
    hdr(ws, r + 1, ["Funcionário", "Função", "Diária (R$)", "Diárias no mês", "", "", "", "Total (R$)"])
    ws.merge_cells(f"D{r+1}:G{r+1}")
    r += 2
    r0 = r
    for n in ORDEM:
        dm = FUNC_MES[n][key]
        if dm == 0:
            continue
        celf(ws, f"A{r}", n)
        celf(ws, f"B{r}", FUNCAO[n])
        celf(ws, f"C{r}", DIARIA[n], MOEDA)
        celf(ws, f"D{r}", dm, NUM, center=True)
        ws.merge_cells(f"D{r}:G{r}")
        celf(ws, f"H{r}", f"=C{r}*D{r}", MOEDA, NEG)
        for cc in "EFG":
            ws[f"{cc}{r}"].border = BOX
        r += 1
    ws[f"A{r}"] = "TOTAL — MÃO DE OBRA DO MÊS"
    ws[f"A{r}"].font = NEG
    ws.merge_cells(f"A{r}:C{r}")
    celf(ws, f"D{r}", f"=SUM(D{r0}:D{r-1})", NUM, NEG, center=True)
    ws.merge_cells(f"D{r}:G{r}")
    celf(ws, f"H{r}", f"=SUM(H{r0}:H{r-1})", MOEDA, NEG)
    for cc in "ABCDEFGH":
        ws[f"{cc}{r}"].fill, ws[f"{cc}{r}"].border = F_SUB, BOX

    r += 2
    barra(ws, r, "3. QUENTINHAS DO MÊS — SEPARADO POR VALOR UNITÁRIO (E OUTROS CONSUMOS)", "H")
    hdr(ws, r + 1, ["Valor unitário / item", "Dias", "", "", "", "Quantidade (un.)", "", "Custo (R$)"])
    ws.merge_cells(f"B{r+1}:E{r+1}")
    ws.merge_cells(f"F{r+1}:G{r+1}")
    r += 2
    r0 = r
    for v in VALORES:
        q = Q_MES[v][key]
        if q == 0:
            continue
        celf(ws, f"A{r}", v, MOEDA, NEG)
        celf(ws, f"B{r}", ", ".join(x[0].strftime("%d/%m") for x in POR_MES[key] if x[3] == v and x[2] > 0), font=PEQ)
        ws.merge_cells(f"B{r}:E{r}")
        celf(ws, f"F{r}", q, INT, center=True)
        ws.merge_cells(f"F{r}:G{r}")
        celf(ws, f"H{r}", f"=A{r}*F{r}", MOEDA, NEG)
        for cc in "CDEG":
            ws[f"{cc}{r}"].border = BOX
        r += 1
    r_qfim = r
    for desc, por_mes in EXTRAS_MES.items():
        if key not in por_mes:
            continue
        q, custo = por_mes[key]
        celf(ws, f"A{r}", desc, font=NEG)
        celf(ws, f"B{r}", ", ".join(x[0].strftime("%d/%m") for x in POR_MES[key] if any(e[0] == desc for e in x[4])), font=PEQ)
        ws.merge_cells(f"B{r}:E{r}")
        celf(ws, f"F{r}", q, INT, center=True)
        ws.merge_cells(f"F{r}:G{r}")
        celf(ws, f"H{r}", custo, MOEDA, NEG)
        for cc in "CDEG":
            ws[f"{cc}{r}"].border = BOX
        r += 1
    ws[f"A{r}"] = "TOTAL — ALIMENTAÇÃO E OUTROS DO MÊS" if r > r_qfim else "TOTAL DE QUENTINHAS DO MÊS"
    ws[f"A{r}"].font = NEG
    ws.merge_cells(f"A{r}:E{r}")
    celf(ws, f"F{r}", f"=SUM(F{r0}:F{r_qfim-1})", INT, NEG, center=True)
    ws.merge_cells(f"F{r}:G{r}")
    celf(ws, f"H{r}", f"=SUM(H{r0}:H{r-1})", MOEDA, NEG)
    for cc in "ABCDEFGH":
        ws[f"{cc}{r}"].fill, ws[f"{cc}{r}"].border = F_SUB, BOX

    r += 2
    barra(ws, r, "4. GASTO GERAL DO MÊS", "H")
    st = STATS[key]
    for i, (lab, val, fmt) in enumerate([("Dias com lançamento", st["dias"], INT),
                                         ("Diárias trabalhadas", st["diarias"], NUM),
                                         ("Total de mão de obra", st["mo"], MOEDA),
                                         ("Total de alimentação e outros", st["alim"], MOEDA)]):
        rr = r + 1 + i
        celf(ws, f"A{rr}", lab)
        ws.merge_cells(f"A{rr}:G{rr}")
        celf(ws, f"H{rr}", val, fmt)
        for cc in "BCDEFG":
            ws[f"{cc}{rr}"].border = BOX
    rr = r + 5
    ws[f"A{rr}"] = f"GASTO GERAL — {nome_up}"
    ws[f"A{rr}"].font = Font(name=A, size=12, bold=True)
    ws.merge_cells(f"A{rr}:G{rr}")
    ws[f"H{rr}"] = f"=H{rr-2}+H{rr-1}"
    ws[f"H{rr}"].number_format, ws[f"H{rr}"].font = MOEDA, Font(name=A, size=12, bold=True)
    for cc in "ABCDEFGH":
        ws[f"{cc}{rr}"].fill, ws[f"{cc}{rr}"].border = F_TOT, BOXM

    for cc, w in {"A": 26, "B": 13, "C": 13, "D": 12, "E": 17, "F": 15, "G": 15, "H": 14, "I": 18}.items():
        ws.column_dimensions[cc].width = w
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

# --- Resumo Mensal
nm = len(MESES)
c_tot, c_val = CL(4 + nm), CL(5 + nm)
ULT = c_val
rs = wb.create_sheet("Resumo Mensal", 0)
rs["A1"] = "RESUMO MENSAL DA OBRA"
rs["A1"].font = Font(name=A, size=14, bold=True, color="1F3864")
rs.merge_cells(f"A1:{ULT}1")
rs["A2"] = OBRA
rs["A2"].font = SUB
rs.merge_cells(f"A2:{ULT}2")

barra(rs, 4, "1. FECHAMENTO POR MÊS", "H")
hdr(rs, 5, ["Mês", "Período", "Dias", "Diárias", "Mão de Obra (R$)", "Quentinhas (un.)",
            "Alim. e outros (R$)", "Total (R$)"])
r = 6
for key, (nome_up, nome, periodo, parcial) in MESES.items():
    st = STATS[key]
    celf(rs, f"A{r}", nome + (" (parcial)" if parcial else ""))
    celf(rs, f"B{r}", periodo)
    celf(rs, f"C{r}", st["dias"], INT, center=True)
    celf(rs, f"D{r}", st["diarias"], NUM, center=True)
    celf(rs, f"E{r}", st["mo"], MOEDA)
    celf(rs, f"F{r}", st["quent"], INT, center=True)
    celf(rs, f"G{r}", st["alim"], MOEDA)
    celf(rs, f"H{r}", f"=E{r}+G{r}", MOEDA, NEG)
    r += 1
rs[f"A{r}"] = "TOTAL ACUMULADO DA OBRA"
rs[f"A{r}"].font = NEG
rs.merge_cells(f"A{r}:B{r}")
for col, fmt in [("C", INT), ("D", NUM), ("E", MOEDA), ("F", INT), ("G", MOEDA), ("H", MOEDA)]:
    celf(rs, f"{col}{r}", f"=SUM({col}6:{col}{r-1})", fmt, NEG, center=col in "CDF")
for cc in "ABCDEFGH":
    rs[f"{cc}{r}"].fill, rs[f"{cc}{r}"].border = F_TOT, BOX

r += 2
barra(rs, r, "2. DIÁRIAS POR FUNCIONÁRIO, MÊS A MÊS", ULT)
cols_m = [MESES[k][1] for k in MESES]
hdr(rs, r + 1, ["Funcionário", "Função", "Diária (R$)"] + cols_m + ["Total diárias", "Total pago (R$)"])
r += 2
r0 = r
for n in FUNC_ATIVOS:
    celf(rs, f"A{r}", n)
    celf(rs, f"B{r}", FUNCAO[n])
    celf(rs, f"C{r}", DIARIA[n], MOEDA)
    for jx, k in enumerate(MESES):
        celf(rs, f"{CL(4+jx)}{r}", FUNC_MES[n][k], NUM, center=True)
    celf(rs, f"{c_tot}{r}", f"=SUM(D{r}:{CL(3+nm)}{r})", NUM, center=True)
    celf(rs, f"{c_val}{r}", f"=C{r}*{c_tot}{r}", MOEDA, NEG)
    r += 1
rs[f"A{r}"] = "TOTAL"
rs[f"A{r}"].font = NEG
rs.merge_cells(f"A{r}:C{r}")
for ci in range(4, 5 + nm):
    celf(rs, f"{CL(ci)}{r}", f"=SUM({CL(ci)}{r0}:{CL(ci)}{r-1})", NUM, NEG, center=True)
celf(rs, f"{c_val}{r}", f"=SUM({c_val}{r0}:{c_val}{r-1})", MOEDA, NEG)
for ci in range(1, 6 + nm):
    rs.cell(row=r, column=ci).fill = F_SUB
    rs.cell(row=r, column=ci).border = BOX

r += 2
barra(rs, r, "3. QUENTINHAS POR VALOR UNITÁRIO, MÊS A MÊS (UNIDADES) — E OUTROS CONSUMOS", ULT)
hdr(rs, r + 1, ["Valor unitário / item", "", ""] + cols_m + ["Total (un.)", "Custo (R$)"])
rs.merge_cells(f"A{r+1}:C{r+1}")
r += 2
r0 = r
for v in VALORES:
    celf(rs, f"A{r}", v, MOEDA, NEG)
    rs.merge_cells(f"A{r}:C{r}")
    for jx, k in enumerate(MESES):
        celf(rs, f"{CL(4+jx)}{r}", Q_MES[v][k], INT, center=True)
    celf(rs, f"{c_tot}{r}", f"=SUM(D{r}:{CL(3+nm)}{r})", INT, center=True)
    celf(rs, f"{c_val}{r}", f"=A{r}*{c_tot}{r}", MOEDA, NEG)
    for cc in "BC":
        rs[f"{cc}{r}"].border = BOX
    r += 1
r_qfim = r
for desc, por_mes in EXTRAS_MES.items():
    celf(rs, f"A{r}", desc, font=NEG)
    rs.merge_cells(f"A{r}:C{r}")
    tq, tc = 0, 0.0
    for jx, k in enumerate(MESES):
        q, custo = por_mes.get(k, (0, 0.0))
        tq += q
        tc += custo
        celf(rs, f"{CL(4+jx)}{r}", q, INT, center=True)
    celf(rs, f"{c_tot}{r}", tq, INT, center=True)
    celf(rs, f"{c_val}{r}", tc, MOEDA, NEG)
    for cc in "BC":
        rs[f"{cc}{r}"].border = BOX
    r += 1
rs[f"A{r}"] = "TOTAL — ALIMENTAÇÃO E OUTROS" if EXTRAS_MES else "TOTAL"
rs[f"A{r}"].font = NEG
rs.merge_cells(f"A{r}:C{r}")
for ci in range(4, 5 + nm):
    celf(rs, f"{CL(ci)}{r}", f"=SUM({CL(ci)}{r0}:{CL(ci)}{r_qfim-1})", INT, NEG, center=True)
celf(rs, f"{c_val}{r}", f"=SUM({c_val}{r0}:{c_val}{r-1})", MOEDA, NEG)
for ci in range(1, 6 + nm):
    rs.cell(row=r, column=ci).fill = F_SUB
    rs.cell(row=r, column=ci).border = BOX

r += 2
rs[f"A{r}"] = ("Nota: o primeiro mês começa no início da obra. O mês corrente está parcial — inclui apenas os dias "
               "já lançados. Semanas que cruzam a virada do mês ficam divididas: no relatório semanal aparecem "
               "inteiras; aqui cada dia conta no seu mês do calendário. Na tabela 3, a coluna de unidades do "
               "total soma só quentinhas.")
rs[f"A{r}"].font = PEQ
rs.merge_cells(f"A{r}:{ULT}{r+1}")
rs[f"A{r}"].alignment = Alignment(wrap_text=True, vertical="top")
for cc, w in {"A": 22, "B": 22, "C": 13}.items():
    rs.column_dimensions[cc].width = w
for ci in range(4, 6 + nm):
    rs.column_dimensions[CL(ci)].width = 16
rs.sheet_view.showGridLines = False
rs.page_setup.orientation = "portrait"
rs.page_setup.fitToWidth = rs.page_setup.fitToHeight = 1
rs.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(os.path.join(ARGS.saida, f"Fechamento_Mensal_{SLUG}.xlsx"))

# ================================================================== PDF
W, H = landscape(A4)
M = 26
AZUL = HexColor("#1F3864")
AZUL_CL = HexColor("#D9E1F2")
CINZA = HexColor("#F2F2F2")
LINHA = HexColor("#BFBFBF")
VERDE = HexColor("#C6E0B4")
TXT = HexColor("#262626")
CZ = HexColor("#595959")

c = canvas.Canvas(os.path.join(ARGS.saida, f"Fechamento_Mensal_{SLUG}.pdf"), pagesize=landscape(A4))
c.setTitle(f"Fechamento Mensal — {OBRA_J['nome']}")
y = H - M
c.setFillColor(AZUL)
c.rect(M, y - 34, W - 2 * M, 34, stroke=0, fill=1)
c.setFillColor(white)
c.setFont("Helvetica-Bold", 14)
c.drawString(M + 10, y - 15, f"FECHAMENTO MENSAL DE DIÁRIAS E ALIMENTAÇÃO — "
                             f"{list(MESES.values())[0][0]} A {list(MESES.values())[-1][0]}")
c.setFont("Helvetica", 8.5)
c.drawString(M + 10, y - 27, f"Obra: {OBRA_J['nome']}  ·  {OBRA_J['endereco']}")
c.setFont("Helvetica-Bold", 10)
c.drawRightString(W - M - 10, y - 17, OBRA_J["empresa"].upper())
top = y - 34 - 10
GAP = 12
NCOL = max(3, nm)
CW = (W - 2 * M - (NCOL - 1) * GAP) / NCOL
BH = 168


def card_mes(x, ytop, key):
    nome_up, nome, periodo, parcial = MESES[key]
    st = STATS[key]
    c.setStrokeColor(AZUL)
    c.setLineWidth(0.8)
    c.rect(x, ytop - BH, CW, BH, stroke=1, fill=0)
    c.setFillColor(AZUL)
    c.rect(x, ytop - 16, CW, 16, stroke=0, fill=1)
    c.setFillColor(white)
    c.setFont("Helvetica-Bold", 9)
    c.drawString(x + 6, ytop - 11.5, nome_up + ("  ·  PARCIAL" if parcial else ""))
    c.setFont("Helvetica", 7.5)
    c.drawRightString(x + CW - 6, ytop - 11.5, periodo.replace(" — parcial", ""))
    yy = ytop - 16
    c.setFillColor(AZUL_CL)
    c.rect(x, yy - 11, CW, 11, stroke=0, fill=1)
    c.setFillColor(AZUL)
    c.setFont("Helvetica-Bold", 6.5)
    c.drawString(x + 5, yy - 8, "SEMANA")
    c.drawCentredString(x + 0.50 * CW, yy - 8, "DIAS")
    c.drawCentredString(x + 0.64 * CW, yy - 8, "DIÁRIAS")
    c.drawRightString(x + CW - 5, yy - 8, "TOTAL")
    yy -= 11
    for n, ini, fim, s in SEM_MES[key]:
        c.setFillColor(TXT)
        c.setFont("Helvetica", 7)
        c.drawString(x + 5, yy - 8, f"Sem. {n}  ({ini.strftime('%d/%m')} a {fim.strftime('%d/%m')})")
        c.drawCentredString(x + 0.50 * CW, yy - 8, str(s["dias"]))
        c.drawCentredString(x + 0.64 * CW, yy - 8, fmtn(s["diarias"]))
        c.drawRightString(x + CW - 5, yy - 8, brl(s["total"]))
        c.setStrokeColor(LINHA)
        c.setLineWidth(0.3)
        c.line(x + 4, yy - 11, x + CW - 4, yy - 11)
        yy -= 11
    linhas = [("Quentinhas", f"{st['quent']} un.", brl(st["quent_custo"])),
              ("Mão de obra", f"{fmtn(st['diarias'])} diária" + ("s" if st["diarias"] != 1 else ""), brl(st["mo"])),
              ("Dias com lançamento", "", str(st["dias"]))]
    if st["extras"]:
        linhas.insert(0, ("Outros consumos", "gelo, entregas…", brl(st["extras"])))
    yb0 = ytop - BH + 17
    for i, (lab, mid, val) in enumerate(linhas):
        yb = yb0 + i * 12
        c.setFillColor(CINZA if i % 2 == 0 else white)
        c.rect(x, yb, CW, 12, stroke=0, fill=1)
        c.setFillColor(TXT)
        c.setFont("Helvetica-Bold", 7)
        c.drawString(x + 5, yb + 3.5, lab)
        c.setFont("Helvetica", 7)
        c.drawCentredString(x + 0.60 * CW, yb + 3.5, mid)
        c.setFont("Helvetica-Bold", 7)
        c.drawRightString(x + CW - 5, yb + 3.5, val)
    c.setFillColor(VERDE)
    c.rect(x, ytop - BH, CW, 17, stroke=0, fill=1)
    c.setStrokeColor(AZUL)
    c.setLineWidth(0.8)
    c.rect(x, ytop - BH, CW, 17, stroke=1, fill=0)
    c.setFillColor(AZUL)
    c.setFont("Helvetica-Bold", 9)
    c.drawString(x + 6, ytop - BH + 5.5, f"TOTAL DE {nome_up.split('/')[0]}")
    c.drawRightString(x + CW - 6, ytop - BH + 5.5, brl(st["total"]))


for i, key in enumerate(MESES):
    card_mes(M + i * (CW + GAP), top, key)

ry = top - BH - 14
c.setFillColor(AZUL)
c.rect(M, ry - 15, W - 2 * M, 15, stroke=0, fill=1)
c.setFillColor(white)
c.setFont("Helvetica-Bold", 10)
c.drawString(M + 8, ry - 11, "DETALHAMENTO MÊS A MÊS")
ry -= 15
LW = (W - 2 * M) * 0.58
RW = (W - 2 * M) - LW - GAP
RX = M + LW + GAP
ABREV = [MESES[k][1][:3].upper() for k in MESES]
ks = list(MESES)
POS_M = [0.50 + i * (0.30 / max(nm, 1)) for i in range(nm)]
POS_Q = [0.34 + i * (0.30 / max(nm, 1)) for i in range(nm)]

# esquerda: por funcionário
BLH = 13 + 11 + len(FUNC_ATIVOS) * 11 + 14
c.setStrokeColor(AZUL)
c.setLineWidth(0.8)
c.rect(M, ry - BLH, LW, BLH, stroke=1, fill=0)
c.setFillColor(AZUL_CL)
c.rect(M, ry - 13, LW, 13, stroke=0, fill=1)
c.setFillColor(AZUL)
c.setFont("Helvetica-Bold", 8)
c.drawString(M + 6, ry - 9, "DIÁRIAS POR FUNCIONÁRIO, MÊS A MÊS")
yy = ry - 13
c.setFillColor(CINZA)
c.rect(M, yy - 11, LW, 11, stroke=0, fill=1)
c.setFillColor(AZUL)
c.setFont("Helvetica-Bold", 6.5)
c.drawString(M + 6, yy - 7.8, "FUNCIONÁRIO")
c.drawString(M + 0.24 * LW, yy - 7.8, "FUNÇÃO")
c.drawCentredString(M + 0.42 * LW, yy - 7.8, "DIÁRIA")
for jx, xk in enumerate(POS_M):
    c.drawCentredString(M + xk * LW, yy - 7.8, ABREV[jx])
c.drawCentredString(M + 0.80 * LW, yy - 7.8, "TOTAL")
c.drawRightString(M + LW - 6, yy - 7.8, "TOTAL PAGO")
yy -= 11
tot_pago = 0
for n in FUNC_ATIVOS:
    fm = FUNC_MES[n]
    t = sum(fm.values())
    v = t * DIARIA[n]
    tot_pago += v
    c.setFillColor(TXT)
    c.setFont("Helvetica", 7)
    c.drawString(M + 6, yy - 8, n)
    c.setFont("Helvetica", 6.5)
    c.drawString(M + 0.24 * LW, yy - 8, FUNCAO[n])
    c.setFont("Helvetica", 7)
    c.drawCentredString(M + 0.42 * LW, yy - 8, brl(DIARIA[n]))
    for jx, xk in enumerate(POS_M):
        val = fm[ks[jx]]
        c.setFillColor(TXT if val else CZ)
        c.drawCentredString(M + xk * LW, yy - 8, fmtn(val) if val else "–")
    c.setFillColor(TXT)
    c.setFont("Helvetica-Bold", 7)
    c.drawCentredString(M + 0.80 * LW, yy - 8, fmtn(t))
    c.drawRightString(M + LW - 6, yy - 8, brl(v))
    c.setStrokeColor(LINHA)
    c.setLineWidth(0.3)
    c.line(M + 4, yy - 11, M + LW - 4, yy - 11)
    yy -= 11
c.setFillColor(CINZA)
c.rect(M, yy - 14, LW, 14, stroke=0, fill=1)
c.setFillColor(AZUL)
c.setFont("Helvetica-Bold", 7.5)
c.drawString(M + 6, yy - 9.5, "TOTAL — MÃO DE OBRA")
for jx, xk in enumerate(POS_M):
    c.drawCentredString(M + xk * LW, yy - 9.5, fmtn(STATS[ks[jx]]["diarias"]))
c.drawCentredString(M + 0.80 * LW, yy - 9.5, fmtn(TOTAL["diarias"]))
c.drawRightString(M + LW - 6, yy - 9.5, brl(tot_pago))

# direita: quentinhas por valor por mês + extras
BQ = 13 + 11 + (len(VALORES) + len(EXTRAS_MES)) * 12 + 14
c.setStrokeColor(AZUL)
c.setLineWidth(0.8)
c.rect(RX, ry - BQ, RW, BQ, stroke=1, fill=0)
c.setFillColor(AZUL_CL)
c.rect(RX, ry - 13, RW, 13, stroke=0, fill=1)
c.setFillColor(AZUL)
c.setFont("Helvetica-Bold", 8)
c.drawString(RX + 6, ry - 9, "QUENTINHAS POR VALOR, MÊS A MÊS" + (" · E OUTROS" if EXTRAS_MES else ""))
yy = ry - 13
c.setFillColor(CINZA)
c.rect(RX, yy - 11, RW, 11, stroke=0, fill=1)
c.setFillColor(AZUL)
c.setFont("Helvetica-Bold", 6.5)
c.drawString(RX + 6, yy - 7.5, "VALOR")
for jx, xk in enumerate(POS_Q):
    c.drawCentredString(RX + xk * RW, yy - 7.5, ABREV[jx])
c.drawCentredString(RX + 0.73 * RW, yy - 7.5, "TOTAL")
c.drawRightString(RX + RW - 6, yy - 7.5, "CUSTO")
yy -= 11


def linha_q(lab, por_mes_fn, total_qtd, custo):
    global yy
    c.setFillColor(TXT)
    c.setFont("Helvetica-Bold", 7.5)
    c.drawString(RX + 6, yy - 8.5, lab)
    c.setFont("Helvetica", 7)
    for jx, xk in enumerate(POS_Q):
        val = por_mes_fn(ks[jx])
        c.setFillColor(TXT if val else CZ)
        c.drawCentredString(RX + xk * RW, yy - 8.5, str(val) if val else "–")
    c.setFillColor(TXT)
    c.setFont("Helvetica-Bold", 7)
    c.drawCentredString(RX + 0.73 * RW, yy - 8.5, f"{total_qtd} un.")
    c.drawRightString(RX + RW - 6, yy - 8.5, brl(custo))
    c.setStrokeColor(LINHA)
    c.setLineWidth(0.3)
    c.line(RX + 4, yy - 12, RX + RW - 4, yy - 12)
    yy -= 12


for v in VALORES:
    t = sum(Q_MES[v].values())
    linha_q(brl(v), lambda k, v=v: Q_MES[v][k], t, t * v)
for desc, por_mes in EXTRAS_MES.items():
    tq = sum(q for q, _ in por_mes.values())
    tc = sum(cst for _, cst in por_mes.values())
    linha_q(desc, lambda k, pm=por_mes: pm.get(k, (0, 0))[0], tq, tc)
c.setFillColor(CINZA)
c.rect(RX, yy - 14, RW, 14, stroke=0, fill=1)
c.setFillColor(AZUL)
c.setFont("Helvetica-Bold", 7.5)
c.drawString(RX + 6, yy - 9.5, "TOTAL — ALIMENTAÇÃO E OUTROS" if EXTRAS_MES else "TOTAL — ALIMENTAÇÃO")
c.drawRightString(RX + RW - 6, yy - 9.5, brl(TOTAL["alim"]))

yy = ry - BQ - 8
c.setFillColor(TXT)
c.setFont("Helvetica", 8)
c.drawString(RX + 6, yy - 9, "Mão de obra — todos os meses")
c.setFont("Helvetica-Bold", 8)
c.drawRightString(RX + RW - 6, yy - 9, brl(TOTAL["mo"]))
c.setFont("Helvetica", 8)
c.drawString(RX + 6, yy - 20, "Alimentação e outros — todos os meses")
c.setFont("Helvetica-Bold", 8)
c.drawRightString(RX + RW - 6, yy - 20, brl(TOTAL["alim"]))
c.setStrokeColor(LINHA)
c.setLineWidth(0.4)
c.line(RX + 4, yy - 24, RX + RW - 4, yy - 24)
c.setFillColor(VERDE)
c.rect(RX, yy - 47, RW, 19, stroke=0, fill=1)
c.setStrokeColor(AZUL)
c.setLineWidth(1.2)
c.rect(RX, yy - 47, RW, 19, stroke=1, fill=0)
c.setFillColor(AZUL)
c.setFont("Helvetica-Bold", 11)
c.drawString(RX + 8, yy - 41, "TOTAL ACUMULADO DA OBRA")
c.drawRightString(RX + RW - 8, yy - 41, brl(TOTAL["total"]))

yy = ry - BLH - 8
c.setFillColor(CZ)
c.setFont("Helvetica-Oblique", 6.8)
c.drawString(M + 4, yy - 8, "O primeiro mês começa no início da obra. O mês corrente é parcial: inclui apenas os dias já lançados.")
c.drawString(M + 4, yy - 17, "Semanas que cruzam a virada do mês ficam divididas: cada dia conta no seu mês do calendário.")
c.drawString(M + 4, yy - 26, "Diária cheia por função conforme cadastro da equipe; meia diária = fator 0,5. Outros consumos: gelo, entrega de quentinha.")

c.setFillColor(HexColor("#808080"))
c.setFont("Helvetica-Oblique", 6.5)
c.drawString(M, M - 12, f"{OBRA_J['empresa']} · Resp. téc. {OBRA_J['responsavel']} · "
                        f"Valores de diária e quentinha conforme informado pela obra · Emitido em {HOJE.strftime('%d/%m/%Y')}")
c.drawRightString(W - M, M - 12, "Página 1 de 1")
c.showPage()
c.save()

for k in MESES:
    print(MESES[k][1], STATS[k])
print("TOTAL", TOTAL)
